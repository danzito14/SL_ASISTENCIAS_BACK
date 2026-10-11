# app/services/reporte_service.py
"""
Generación de reportes descargables (CSV / XLSX) de las distintas entidades.
Cada reporte se acota por la empresa del usuario (None = super-admin → todas).
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from uuid import UUID
from zoneinfo import ZoneInfo

from fastapi import HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func, text
from sqlalchemy.orm import Session

from src.models.AreaTrabajo_Trabajador_Model import AreaTrabajo, Trabajador
from src.models.Asistencia_HistorialAuditoria_Model import Asistencia
from src.models.Embedding_Model import Embedding
from src.models.Empresa_Model import Empresa
from src.models.Incidencia_Model import Incidencia
from src.models.IntentoAcceso_Model import IntentoAcceso
from src.models.Puerta_Model import PuertaAcceso

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Diseño del XLSX: fila 1 = logo (A1) + título (C1:D1) + rango (E1);
# fila 2 = encabezados de una tabla de Excel; datos desde la fila 3.
FILA_TITULO = 1
FILA_ENCABEZADOS = 2
ALTO_TITULO_PT = 67.8
ANCHO_COL_LOGO = 24.9          # ancho de la columna A (caracteres)
CAJA_LOGO_PX = (170, 82)       # el logo entra completo en esta caja (contain)
ESTILO_TABLA = "TableStyleLight9"
# Logos por empresa: assets/logos/logo_empresa_<id_empresa>.<ext> (mismos ids que
# la tabla empresas). Sin logo propio (o "todas las empresas") → el de la empresa 1.
LOGOS_DIR = Path(__file__).resolve().parents[2] / "assets" / "logos"
LOGO_EXTS = (".png", ".jpg", ".jpeg", ".webp")
ID_EMPRESA_LOGO_DEFAULT = 1
_logos_cache: dict[int | None, tuple[bytes, int, int] | None] = {}


def _ruta_logo(id_empresa: int) -> Path | None:
    for ext in LOGO_EXTS:
        ruta = LOGOS_DIR / f"logo_empresa_{id_empresa}{ext}"
        if ruta.is_file():
            return ruta
    return None


def _logo_empresa(id_empresa: int | None) -> tuple[bytes, int, int] | None:
    """PNG del logo reducido a la caja del título → (bytes, ancho_px, alto_px)."""
    if id_empresa in _logos_cache:
        return _logos_cache[id_empresa]
    res = None
    ruta = (_ruta_logo(id_empresa) if id_empresa is not None else None) \
        or _ruta_logo(ID_EMPRESA_LOGO_DEFAULT)
    if ruta is not None:
        try:
            from PIL import Image as PILImage
            with PILImage.open(ruta) as im:
                im = im.convert("RGBA")
                esc = min(CAJA_LOGO_PX[0] / im.width, CAJA_LOGO_PX[1] / im.height)
                ancho, alto = max(1, round(im.width * esc)), max(1, round(im.height * esc))
                # Se guarda al doble de resolución (nitidez) pero sin el peso del original.
                im.thumbnail((ancho * 2, alto * 2))
                buf = io.BytesIO()
                im.save(buf, format="PNG", optimize=True)
                res = (buf.getvalue(), ancho, alto)
        except Exception:
            res = None   # imagen inválida o sin Pillow: el reporte sale sin logo
    _logos_cache[id_empresa] = res
    return res


_MESES = ("Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")
# Razón social al final del nombre ("SA de CV", "S.A.", "S de RL de CV"...): no va en el archivo.
_RAZON_SOCIAL = re.compile(
    r"\s+S\.?\s?(A|DE\s+R\.?\s?L)\.?(\s+DE\s+C\.?\s?V\.?)?\s*$", re.IGNORECASE)


def _slug_archivo(texto: str) -> str:
    """'SL Agricola SA de CV' → 'SL_Agricola' (sin acentos ni símbolos, apto para nombre de archivo)."""
    texto = _RAZON_SOCIAL.sub("", texto.strip())
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^A-Za-z0-9]+", " ", texto)
    return "_".join(texto.split())


def _fecha_archivo(d: date) -> str:
    """date(2026, 10, 7) → '07Oct2026'."""
    return f"{d.day:02d}{_MESES[d.month - 1]}{d.year}"


def _celda(v):
    """Normaliza un valor para exportar (fechas a texto, Decimal a float, UUID a texto)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, UUID):   # openpyxl no acepta objetos UUID; las PKs ahora son UUID
        return str(v)
    return v


def _a_local(dt: datetime, zona: str | None, cache: dict) -> datetime:
    """Pasa un timestamptz a la hora local de la empresa (UTC si no hay zona válida)."""
    if zona not in cache:
        try:
            cache[zona] = ZoneInfo(zona) if zona else timezone.utc
        except Exception:
            cache[zona] = timezone.utc
    # fecha_hora es timestamptz (aware); si viniera naive, se asume UTC.
    aware = dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    return aware.astimezone(cache[zona])


def _si_no(v: bool | None) -> str:
    return "" if v is None else ("Sí" if v else "No")


def _hasta_fin_de_dia(fecha_fin: date) -> datetime:
    """Convierte una fecha 'fin' (inclusiva) al límite superior para timestamps."""
    return datetime.combine(fecha_fin + timedelta(days=1), datetime.min.time())


class ReporteService:

    # ── Título del XLSX ────────────────────────────────────────────────────────
    def _titulo(
        self, db: Session, id_empresa: int | None, titulo: str, tipo_archivo: str,
        fecha_inicio: date | None = None, fecha_fin: date | None = None,
    ) -> dict:
        """
        Datos del título: "<TITULO> DE <EMPRESA>" + "DEL <inicio> AL <fin>".
        Sin fecha de inicio el rango queda "AL <fin>" (todo el histórico hasta 'fin');
        sin fecha fin se usa el día de hoy.

        También el nombre del archivo: Reporte_<tipo_archivo>_<Empresa>_<07Oct2026>
        (fecha de descarga en la zona horaria de la empresa).
        """
        if id_empresa is None:
            empresa = "Todas las empresas"
        else:
            fila = db.query(Empresa.nombre_empresa).filter(Empresa.id_empresa == id_empresa).first()
            empresa = fila[0] if fila and fila[0] else f"Empresa {id_empresa}"
        hoy = datetime.now(self._tz_empresa(db, id_empresa)).date()
        fin = (fecha_fin or hoy).isoformat()
        rango = f"DEL {fecha_inicio.isoformat()} AL {fin}" if fecha_inicio else f"AL {fin}"
        return {
            "texto": f"{titulo.upper()} DE {empresa.upper()}",
            "rango": rango,
            "logo": _logo_empresa(id_empresa),
            "archivo": f"Reporte_{tipo_archivo}_{_slug_archivo(empresa)}_{_fecha_archivo(hoy)}",
        }

    def _escribir_titulo(self, ws, titulo: dict) -> None:
        """Fila 1: logo (A1), título (C1:D1 combinadas) y rango de fechas (E1)."""
        from openpyxl.styles import Alignment, Font

        ws.row_dimensions[FILA_TITULO].height = ALTO_TITULO_PT
        fuente = Font(bold=True, size=16)
        c = ws.cell(row=FILA_TITULO, column=3, value=titulo["texto"])
        c.font = fuente
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=FILA_TITULO, start_column=3, end_row=FILA_TITULO, end_column=4)
        c = ws.cell(row=FILA_TITULO, column=5, value=titulo["rango"])
        c.font = fuente
        c.alignment = Alignment(vertical="center")

        if titulo["logo"] is None:
            return
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU

        datos, ancho, alto = titulo["logo"]
        img = Image(io.BytesIO(datos))
        img.width, img.height = ancho, alto
        # Centrado dentro de A1 (ancho de columna en px ≈ caracteres * 7 + 5).
        caja_w = int(ANCHO_COL_LOGO * 7 + 5)
        caja_h = int(ALTO_TITULO_PT * 96 / 72)
        marca = AnchorMarker(
            col=0, colOff=pixels_to_EMU(max(0, (caja_w - ancho) // 2)),
            row=FILA_TITULO - 1, rowOff=pixels_to_EMU(max(0, (caja_h - alto) // 2)),
        )
        img.anchor = OneCellAnchor(
            _from=marca, ext=XDRPositiveSize2D(pixels_to_EMU(ancho), pixels_to_EMU(alto)))
        ws.add_image(img)

    # ── Exportador genérico ────────────────────────────────────────────────────
    def _exportar(
        self,
        columnas: list[tuple[str, str]],   # [(clave, "Encabezado"), ...]
        filas: list[dict],
        nombre_base: str,
        formato: str,
        titulo: dict | None = None,        # ver _titulo(); solo aplica al XLSX
        colores: dict[str, dict[str, str]] | None = None,  # {clave: {valor: "RRGGBB"}} (XLSX)
    ) -> StreamingResponse:
        claves = [c[0] for c in columnas]
        encabezados = [c[1] for c in columnas]

        if formato == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(encabezados)
            for f in filas:
                w.writerow([_celda(f.get(k)) for k in claves])
            # utf-8-sig (BOM) para que Excel respete acentos al abrir el CSV.
            data = buf.getvalue().encode("utf-8-sig")
            media = "text/csv; charset=utf-8"

        elif formato == "xlsx":
            try:
                from openpyxl import Workbook
            except ImportError:
                raise HTTPException(
                    status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                    detail="Falta la librería 'openpyxl' para exportar a XLSX (pip install openpyxl).",
                )
            from openpyxl.styles import PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo

            wb = Workbook()
            ws = wb.active
            ws.title = nombre_base[:31]
            if titulo is not None:
                self._escribir_titulo(ws, titulo)
                fila_enc = FILA_ENCABEZADOS
            else:
                fila_enc = 1
            for col, enc in enumerate(encabezados, start=1):
                ws.cell(row=fila_enc, column=col, value=enc)
            rellenos = {
                k: {v: PatternFill("solid", fgColor=hexc) for v, hexc in mapa.items()}
                for k, mapa in (colores or {}).items()
            }
            anchos = [len(str(e)) for e in encabezados]
            for r, f in enumerate(filas, start=fila_enc + 1):
                for col, k in enumerate(claves, start=1):
                    v = _celda(f.get(k))
                    celda = ws.cell(row=r, column=col, value=v)
                    if k in rellenos and v in rellenos[k]:
                        celda.fill = rellenos[k][v]
                    anchos[col - 1] = max(anchos[col - 1], len(str(v)))
            for col, ancho in enumerate(anchos, start=1):
                ancho = min(max(ancho + 2, 8), 50)
                if col == 1 and titulo is not None:
                    ancho = max(ancho, ANCHO_COL_LOGO)   # cabe el logo
                ws.column_dimensions[get_column_letter(col)].width = ancho

            # Tabla de Excel (filtros + filas bandeadas). Necesita al menos una fila
            # de datos: sin datos se deja una vacía bajo los encabezados.
            ultima = fila_enc + max(len(filas), 1)
            tabla = Table(
                displayName="Reporte",
                ref=f"A{fila_enc}:{get_column_letter(len(claves))}{ultima}",
            )
            tabla.tableStyleInfo = TableStyleInfo(
                name=ESTILO_TABLA, showRowStripes=True, showColumnStripes=False,
                showFirstColumn=False, showLastColumn=False,
            )
            ws.add_table(tabla)
            bio = io.BytesIO()
            wb.save(bio)
            data = bio.getvalue()
            media = XLSX_MEDIA

        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El formato debe ser 'csv' o 'xlsx'.",
            )

        nombre = f"{titulo['archivo'] if titulo else nombre_base}.{formato}"
        return StreamingResponse(
            io.BytesIO(data),
            media_type=media,
            headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
        )

    # ── Reportes ───────────────────────────────────────────────────────────────
    def asistencias(
        self, db: Session, id_empresa: int | None, fecha_inicio: date | None,
        fecha_fin: date | None, formato: str, id_trabajador: int | None = None,
        tipo: str | None = None,
    ) -> StreamingResponse:
        """
        Asistencias con el N° de empleado de SYS21 y la fecha/hora en columnas
        separadas, en la zona horaria de la empresa donde se fichó (fecha_hora es UTC).
        'tipo' acota a solo entradas o solo salidas.
        """
        q = (
            db.query(Asistencia, Trabajador, Empresa, PuertaAcceso)
            .join(Trabajador, Trabajador.id_trabajador == Asistencia.id_trabajador)
            .outerjoin(Empresa, Empresa.id_empresa == Asistencia.id_empresa)
            .outerjoin(PuertaAcceso, PuertaAcceso.id_puerta == Asistencia.id_puerta)
        )
        if id_empresa is not None:
            q = q.filter(Asistencia.id_empresa == id_empresa)  # denormalizado (access)
        if id_trabajador is not None:
            q = q.filter(Asistencia.id_trabajador == id_trabajador)
        if tipo is not None:
            q = q.filter(Asistencia.tipo_registro == tipo)
        if fecha_inicio is not None:
            q = q.filter(Asistencia.fecha_hora >= fecha_inicio)
        if fecha_fin is not None:
            q = q.filter(Asistencia.fecha_hora < _hasta_fin_de_dia(fecha_fin))

        zonas: dict = {}
        filas = []
        for a, t, e, p in q.order_by(Asistencia.fecha_hora.desc()).all():
            local = _a_local(a.fecha_hora, e.zona_horaria if e else None, zonas)
            filas.append({
                "id_emp": t.id_emp,
                "origen_nomina": t.origen_nomina,
                "trabajador": f"{t.nombre} {t.apellido}",
                "empresa": e.nombre_empresa if e else None,
                "puerta": p.nombre_puerta if p else None,
                "tipo": a.tipo_registro,
                "fecha": local.date(),
                "hora": local.strftime("%H:%M:%S"),
            })
        columnas = [
            ("id_emp", "N° empleado"), ("origen_nomina", "Nómina"),
            ("trabajador", "Trabajador"), ("empresa", "Empresa"), ("puerta", "Puerta"),
            ("tipo", "Tipo"), ("fecha", "Fecha"), ("hora", "Hora"),
        ]
        nombre = {"entrada": "entradas", "salida": "salidas"}.get(tipo, "asistencias")
        titulo = self._titulo(db, id_empresa, nombre, tipo or "asistencia", fecha_inicio, fecha_fin)
        colores = {"tipo": {"entrada": "558ED5", "salida": "B3A2C7"}}
        return self._exportar(columnas, filas, nombre, formato, titulo, colores)

    def incidencias(
        self, db: Session, id_empresa: int | None, fecha_inicio: date | None,
        fecha_fin: date | None, formato: str, tipo: str | None = None,
        id_trabajador: int | None = None,
    ) -> StreamingResponse:
        q = db.query(Incidencia, Trabajador).join(
            Trabajador, Trabajador.id_trabajador == Incidencia.id_trabajador
        )
        if id_empresa is not None:
            q = q.filter(Incidencia.id_empresa == id_empresa)  # denormalizado (access)
        if tipo is not None:
            q = q.filter(Incidencia.tipo_incidencia == tipo)
        if id_trabajador is not None:
            q = q.filter(Incidencia.id_trabajador == id_trabajador)
        if fecha_inicio is not None:
            q = q.filter(Incidencia.fecha >= fecha_inicio)
        if fecha_fin is not None:
            q = q.filter(Incidencia.fecha <= fecha_fin)

        filas = [
            {
                "id": i.id_incidencia,
                "trabajador": f"{t.nombre} {t.apellido}",
                "tipo": i.tipo_incidencia,
                "fecha": i.fecha,
                "estado": i.estado,
                "descripcion": i.descripcion,
                # Escaneo que originó la incidencia → el front puede pedir
                # GET /escaneos/{id} para la info completa. NULL si no aplica.
                "escaneo_ref": i.id_escaneo_ref,
            }
            for i, t in q.order_by(Incidencia.fecha.desc(), Incidencia.id_incidencia.desc()).all()
        ]
        columnas = [
            ("id", "ID"), ("trabajador", "Trabajador"), ("tipo", "Tipo"),
            ("fecha", "Fecha"), ("estado", "Estado"), ("descripcion", "Descripción"),
            ("escaneo_ref", "Escaneo (ID)"),
        ]
        nombre_titulo = {"falta": "faltas", "fuera_de_area": "fuera de área"}.get(
            tipo, f"incidencias ({tipo.replace('_', ' ')})" if tipo else "incidencias")
        titulo = self._titulo(db, id_empresa, nombre_titulo, tipo or "incidencia",
                              fecha_inicio, fecha_fin)
        return self._exportar(columnas, filas, "incidencias", formato, titulo)

    def retardos(
        self, db: Session, id_empresa: int | None, fecha_inicio: date | None,
        fecha_fin: date | None, formato: str, id_trabajador: int | None = None,
        tolerancia_min: int = 0,
    ) -> StreamingResponse:
        """
        Reporte CALCULADO de retardos: entradas cuya hora LOCAL supera la hora_entrada
        del área del trabajador (más una tolerancia opcional en minutos), con los
        minutos de retraso. Solo considera áreas con hora_entrada definida.

        La hora se compara en la zona horaria de la empresa (fecha_hora es UTC).
        """
        q = (
            db.query(Asistencia, Trabajador, AreaTrabajo, Empresa)
            .join(Trabajador, Trabajador.id_trabajador == Asistencia.id_trabajador)
            .join(AreaTrabajo, AreaTrabajo.id_area == Trabajador.id_area)
            .join(Empresa, Empresa.id_empresa == Asistencia.id_empresa)
            .filter(Asistencia.tipo_registro == "entrada")
            .filter(AreaTrabajo.hora_entrada.isnot(None))
        )
        if id_empresa is not None:
            q = q.filter(Asistencia.id_empresa == id_empresa)
        if id_trabajador is not None:
            q = q.filter(Asistencia.id_trabajador == id_trabajador)
        if fecha_inicio is not None:
            q = q.filter(Asistencia.fecha_hora >= fecha_inicio)
        if fecha_fin is not None:
            q = q.filter(Asistencia.fecha_hora < _hasta_fin_de_dia(fecha_fin))

        # ASC: la primera fila por (trabajador, día local) es su entrada más temprana;
        # así se cuenta UNA sola vez al día (la hora real de llegada).
        vistos: set = set()
        zonas: dict = {}
        filas = []
        for a, t, ar, e in q.order_by(Asistencia.fecha_hora.asc()).all():
            local = _a_local(a.fecha_hora, e.zona_horaria, zonas)
            clave = (t.id_trabajador, local.date())
            if clave in vistos:
                continue
            vistos.add(clave)
            esperada = ar.hora_entrada
            retraso = (local.hour * 60 + local.minute) - (esperada.hour * 60 + esperada.minute)
            if retraso <= tolerancia_min:
                continue  # a tiempo (o dentro de la tolerancia)
            filas.append({
                "id_trabajador": t.id_trabajador,
                "id_emp": t.id_emp,
                "trabajador": f"{t.nombre} {t.apellido}",
                "area": ar.nombre_area,
                "empresa": e.nombre_empresa,
                "fecha": local.date(),
                "hora_esperada": esperada.strftime("%H:%M"),
                "hora_real": local.strftime("%H:%M"),
                "minutos_retardo": retraso,
            })
        # Más recientes primero para el archivo.
        filas.sort(key=lambda r: (r["fecha"], r["minutos_retardo"]), reverse=True)
        columnas = [
            ("id_trabajador", "ID trabajador"), ("id_emp", "N° empleado"),
            ("trabajador", "Trabajador"), ("area", "Área"), ("empresa", "Empresa"),
            ("fecha", "Fecha"), ("hora_esperada", "Hora esperada"),
            ("hora_real", "Hora real"), ("minutos_retardo", "Minutos de retardo"),
        ]
        titulo = self._titulo(db, id_empresa, "retardos", "retardo", fecha_inicio, fecha_fin)
        return self._exportar(columnas, filas, "retardos", formato, titulo)

    def intentos(
        self, db: Session, id_empresa: int | None, fecha_inicio: date | None,
        fecha_fin: date | None, formato: str,
    ) -> StreamingResponse:
        q = db.query(IntentoAcceso, Trabajador).outerjoin(
            Trabajador, Trabajador.id_trabajador == IntentoAcceso.id_trabajador
        )
        if id_empresa is not None:
            q = q.filter(IntentoAcceso.id_empresa == id_empresa)
        if fecha_inicio is not None:
            q = q.filter(IntentoAcceso.fecha >= fecha_inicio)
        if fecha_fin is not None:
            q = q.filter(IntentoAcceso.fecha < _hasta_fin_de_dia(fecha_fin))

        filas = [
            {
                "id": it.id_intento,
                "tipo": it.tipo,
                "fecha": it.fecha,
                "puerta": it.id_puerta,
                "trabajador": f"{t.nombre} {t.apellido}" if t else "",
                "similitud": it.similitud,
            }
            for it, t in q.order_by(IntentoAcceso.fecha.desc()).all()
        ]
        columnas = [
            ("id", "ID"), ("tipo", "Tipo"), ("fecha", "Fecha y hora"),
            ("puerta", "Puerta"), ("trabajador", "Trabajador (si aplica)"),
            ("similitud", "Similitud"),
        ]
        titulo = self._titulo(db, id_empresa, "intentos de acceso", "intento_acceso",
                              fecha_inicio, fecha_fin)
        return self._exportar(columnas, filas, "intentos_acceso", formato, titulo)

    def trabajadores(
        self, db: Session, id_empresa: int | None, formato: str,
        rostro: str | None = None,
    ) -> StreamingResponse:
        """
        Padrón de trabajadores con su N° de empleado de SYS21 y si tienen rostro
        registrado. 'rostro' = 'con' | 'sin' acota a los que tienen (o no) embedding;
        mismo criterio que el conteo con/sin rostro del dashboard.
        """
        # Los joins a área/empresa se conservan: el reporte muestra sus NOMBRES.
        # El filtro por empresa sí usa el id_empresa denormalizado del trabajador.
        # embeddings es 1-a-1 con trabajadores (UNIQUE) → el outer join no duplica filas.
        q = (
            db.query(Trabajador, AreaTrabajo, Empresa, Embedding)
            .join(AreaTrabajo, AreaTrabajo.id_area == Trabajador.id_area)
            .join(Empresa, Empresa.id_empresa == AreaTrabajo.id_empresa)
            .outerjoin(Embedding, Embedding.id_trabajador == Trabajador.id_trabajador)
        )
        if id_empresa is not None:
            q = q.filter(Trabajador.id_empresa == id_empresa)
        if rostro == "con":
            q = q.filter(Embedding.id_embedding.isnot(None))
        elif rostro == "sin":
            q = q.filter(Embedding.id_embedding.is_(None))

        zonas: dict = {}
        filas = [
            {
                "id": t.id_trabajador,
                "id_emp": t.id_emp,
                "origen_nomina": t.origen_nomina,
                "nombre": t.nombre,
                "apellido": t.apellido,
                "area": a.nombre_area,
                "tipo_area": a.tipo_area,
                "empresa": e.nombre_empresa,
                "estado": t.estado,
                "rostro": _si_no(emb is not None),
                "fecha_rostro": (_a_local(emb.fecha_captura, e.zona_horaria, zonas)
                                 if emb and emb.fecha_captura else None),
                "modelo": emb.modelo_ia if emb else None,
            }
            for t, a, e, emb in q.order_by(Trabajador.id_trabajador).all()
        ]
        columnas = [
            ("id", "ID"), ("id_emp", "N° empleado"), ("origen_nomina", "Nómina"),
            ("nombre", "Nombre"), ("apellido", "Apellido"),
            ("area", "Área"), ("tipo_area", "Tipo de área"), ("empresa", "Empresa"),
            ("estado", "Estado"), ("rostro", "Tiene rostro"),
            ("fecha_rostro", "Rostro registrado el"), ("modelo", "Modelo"),
        ]
        nombre = {"con": "trabajadores_con_rostro", "sin": "trabajadores_sin_rostro"}.get(
            rostro, "trabajadores")
        tipo_archivo = {"con": "trabajador_con_rostro", "sin": "trabajador_sin_rostro"}.get(
            rostro, "trabajador")
        # padrón: sin rango → "AL <hoy>"
        titulo = self._titulo(db, id_empresa, nombre.replace("_", " "), tipo_archivo)
        return self._exportar(columnas, filas, nombre, formato, titulo)

    # ── Dashboard (JSON, no archivo) ───────────────────────────────────────────
    def _tz_empresa(self, db: Session, id_empresa: int | None) -> ZoneInfo:
        if id_empresa is not None:
            fila = db.query(Empresa.zona_horaria).filter(Empresa.id_empresa == id_empresa).first()
            if fila and fila[0]:
                try:
                    return ZoneInfo(fila[0])
                except Exception:
                    pass
        return ZoneInfo("America/Mazatlan")

    def dashboard(self, db: Session, id_empresa: int | None, fecha: date | None = None) -> dict:
        """
        Números del día para el panel: presentes/total (global y por tipo de área),
        ausentes, retardos, intentos, incidencias pendientes, y padrón con/sin rostro.
        'presente' = tiene al menos una ENTRADA en el día (zona horaria de la empresa).
        """
        tz = self._tz_empresa(db, id_empresa)
        dia = fecha or datetime.now(tz).date()
        ini = datetime.combine(dia, time.min, tzinfo=tz).astimezone(timezone.utc)
        fin = ini + timedelta(days=1)

        def _emp(q, col):
            return q.filter(col == id_empresa) if id_empresa is not None else q

        # Total de trabajadores activos por tipo de área.
        q_tot = (
            db.query(AreaTrabajo.tipo_area, func.count(Trabajador.id_trabajador))
            .select_from(AreaTrabajo)
            .join(Trabajador, Trabajador.id_area == AreaTrabajo.id_area)
            .filter(Trabajador.estado == "activo")
        )
        q_tot = _emp(q_tot, Trabajador.id_empresa)
        total_por_tipo = {(t or "sin_tipo"): n for t, n in q_tot.group_by(AreaTrabajo.tipo_area).all()}

        # Presentes hoy (trabajadores DISTINTOS con entrada) por tipo de área.
        q_pre = (
            db.query(AreaTrabajo.tipo_area, func.count(func.distinct(Asistencia.id_trabajador)))
            .select_from(Asistencia)
            .join(Trabajador, Trabajador.id_trabajador == Asistencia.id_trabajador)
            .join(AreaTrabajo, AreaTrabajo.id_area == Trabajador.id_area)
            .filter(Asistencia.tipo_registro == "entrada",
                    Asistencia.fecha_hora >= ini, Asistencia.fecha_hora < fin)
        )
        q_pre = _emp(q_pre, Asistencia.id_empresa)
        pres_por_tipo = {(t or "sin_tipo"): n for t, n in q_pre.group_by(AreaTrabajo.tipo_area).all()}

        tipos = set(total_por_tipo) | set(pres_por_tipo)
        por_tipo = {
            t: {"total": total_por_tipo.get(t, 0), "presentes": pres_por_tipo.get(t, 0)}
            for t in sorted(tipos)
        }
        total = sum(total_por_tipo.values())
        presentes = sum(pres_por_tipo.values())

        # Retardos de hoy: 1ª entrada del día por trabajador vs hora_entrada del área.
        q_ret = (
            db.query(Asistencia.id_trabajador, Asistencia.fecha_hora, AreaTrabajo.hora_entrada)
            .select_from(Asistencia)
            .join(Trabajador, Trabajador.id_trabajador == Asistencia.id_trabajador)
            .join(AreaTrabajo, AreaTrabajo.id_area == Trabajador.id_area)
            .filter(Asistencia.tipo_registro == "entrada", AreaTrabajo.hora_entrada.isnot(None),
                    Asistencia.fecha_hora >= ini, Asistencia.fecha_hora < fin)
        )
        q_ret = _emp(q_ret, Asistencia.id_empresa)
        primeras: dict[int, tuple] = {}
        for id_t, fh, he in q_ret.order_by(Asistencia.fecha_hora.asc()).all():
            if id_t not in primeras:
                primeras[id_t] = (fh, he)
        retardos = 0
        for fh, he in primeras.values():
            local = (fh if fh.tzinfo else fh.replace(tzinfo=timezone.utc)).astimezone(tz)
            if (local.hour * 60 + local.minute) > (he.hour * 60 + he.minute):
                retardos += 1

        # Intentos de hoy.
        q_int = db.query(func.count(IntentoAcceso.id_intento)).filter(
            IntentoAcceso.fecha >= ini, IntentoAcceso.fecha < fin)
        q_int = _emp(q_int, IntentoAcceso.id_empresa)
        intentos = q_int.scalar() or 0

        # Incidencias PENDIENTES (todas las abiertas, no solo hoy).
        q_inc = db.query(func.count(Incidencia.id_incidencia)).filter(Incidencia.estado == "pendiente")
        q_inc = _emp(q_inc, Incidencia.id_empresa)
        incidencias_pend = q_inc.scalar() or 0

        # Padrón con/sin rostro (conteo por SQL crudo sobre embeddings).
        sql_cr = ("SELECT count(DISTINCT e.id_trabajador) FROM embeddings e "
                  "JOIN trabajadores t ON t.id_trabajador = e.id_trabajador "
                  "WHERE t.estado='activo'")
        params_cr: dict = {}
        if id_empresa is not None:
            sql_cr += " AND t.id_empresa = :emp"
            params_cr["emp"] = id_empresa
        con_rostro = db.execute(text(sql_cr), params_cr).scalar() or 0

        return {
            "fecha": dia.isoformat(),
            "empresa": id_empresa,
            "trabajadores": {"total": total, "con_rostro": con_rostro, "sin_rostro": max(0, total - con_rostro)},
            "asistencia": {
                "presentes": presentes,
                "ausentes": max(0, total - presentes),
                "por_tipo": por_tipo,          # {oficina:{total,presentes}, empaque:..., campo:...}
            },
            "retardos_hoy": retardos,
            "intentos_hoy": intentos,
            "incidencias_pendientes": incidencias_pend,
        }


reporte_service = ReporteService()
